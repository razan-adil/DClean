import os
import sys
import tempfile
import unittest.mock
import pytest
import openpyxl
import yaml
from pipeline.contract_generator import (
    AutoContractGenerator,
    BaseNLPAssistant,
    RuleAssistedNLP,
    TransformerNLPAssistant,
    main as cli_main
)
from pipeline.contract_engine import ContractEngine

# =====================================================================
# 1. Fuzzy Typo-Tolerant Resolution Tests (ContractEngine)
# =====================================================================

@pytest.fixture
def fuzzy_engine():
    contract = {
        "dataset_id": "fuzzy_test",
        "quarantine_threshold_pct": 15.0,
        "similarity_threshold": 0.82,
        "canonical_mappings": [
            {
                "target_field": "phone",
                "group": "Identity",
                "type": "phone",
                "aliases": ["رقم الجوال", "Phone Number", "Mobile", "رقم الهاتف"],
                "required": True
            },
            {
                "target_field": "email",
                "group": "Identity",
                "type": "email",
                "aliases": ["البريد الإلكتروني", "Email Address"],
                "required": True
            },
            {
                "target_field": "university",
                "group": "Structured/MetaData",
                "type": "string",
                "aliases": ["اسم الجامعة", "University"],
                "required": False
            },
            {
                "target_field": "full_name_en",
                "group": "Identity",
                "type": "string",
                "aliases": ["Full Name (En)", "English Name"],
                "required": False
            },
            {
                "target_field": "full_name_ar",
                "group": "Identity",
                "type": "string",
                "aliases": ["الاسم الثلاثي", "اسم المشارك"],
                "required": False
            }
        ]
    }
    return ContractEngine(contract)

def test_fuzzy_header_matching_typos(fuzzy_engine):
    typo_headers = [
        ("رقم الجواال", "phone"),
        ("Email Adress", "email"),
        ("Phonne", "phone"),
        ("Universty", "university"),
        ("Full Name (Engish)", "full_name_en")
    ]
    raw_list = [h for h, _ in typo_headers]
    resolved, unmapped = fuzzy_engine.resolve_headers(raw_list)

    for h, expected in typo_headers:
        assert resolved.get(h) == expected, f"Failed to fuzzy-match typo header '{h}' -> {expected}"
    assert len(unmapped) == 0

def test_fuzzy_header_no_false_positives(fuzzy_engine):
    unrelated_headers = ["سؤال عام عن المعسكر", "ما هي الصعوبات التي واجهتك؟", "custom_id_column_xyz"]
    resolved, unmapped = fuzzy_engine.resolve_headers(unrelated_headers)
    assert len(resolved) == 0
    assert len(unmapped) == 3

# =====================================================================
# 2. Deep Sample-Based Heuristics Tests (AutoContractGenerator)
# =====================================================================

def test_deep_sample_inference_saudi_phone():
    res = AutoContractGenerator.classify_column("col_val", sample_values=["0501234567", "0559876543"])
    assert res["target_field"] == "phone"
    assert res["group"] == "Identity"
    assert res["type"] == "phone"
    assert res["detection_source"] == "sample_heuristic"
    assert res["confidence"] >= 0.85

def test_deep_sample_inference_academic_and_personal_email():
    res_acad = AutoContractGenerator.classify_column("user_id_ref", sample_values=["student@uj.edu.sa", "admin@kau.edu.sa"])
    assert res_acad["target_field"] == "academic_email"
    assert res_acad["group"] == "Identity"
    assert res_acad["detection_source"] == "sample_heuristic"
    assert res_acad["confidence"] >= 0.90

    res_gen = AutoContractGenerator.classify_column("contact_detail", sample_values=["sara.ahmed@gmail.com", "fahad@hotmail.com"])
    assert res_gen["target_field"] == "email"
    assert res_gen["group"] == "Identity"
    assert res_gen["detection_source"] == "sample_heuristic"

def test_deep_sample_inference_student_id():
    res = AutoContractGenerator.classify_column("رقم_عشوائي", sample_values=["441002938", "442001928", "443009182"])
    assert res["target_field"] == "student_id"
    assert res["group"] == "Structured/MetaData"
    assert res["detection_source"] == "sample_heuristic"
    assert res["confidence"] >= 0.85

def test_deep_sample_inference_linkedin_url():
    res = AutoContractGenerator.classify_column("social_link", sample_values=["https://www.linkedin.com/in/fahad", "linkedin.com/in/salman"])
    assert res["target_field"] == "linkedin_url"
    assert res["group"] == "Structured/MetaData"
    assert res["detection_source"] == "sample_heuristic"

def test_deep_sample_inference_saudi_academic_level():
    res = AutoContractGenerator.classify_column("معلومات_المرحلة", sample_values=["سنة أولى", "سنة ثالثة", "خريج"])
    assert res["target_field"] == "academic_lvl"
    assert res["group"] == "Structured/MetaData"
    assert res["detection_source"] == "sample_heuristic"

# =====================================================================
# 3. Confidence Scoring & Detection Source
# =====================================================================

def test_confidence_scoring_and_provenance():
    res_pat = AutoContractGenerator.classify_column("الاسم الثلاثي")
    assert res_pat["confidence"] >= 0.90
    assert res_pat["detection_source"] == "header_pattern"

    res_fall = AutoContractGenerator.classify_column("حقل_غير_مفهوم_بتاتاً")
    assert res_fall["confidence"] < 0.60
    assert res_fall["detection_source"] == "fallback"

def test_contract_generation_preserves_confidence():
    headers = ["الاسم الثلاثي", "رقم الجوال", "غير_معروف"]
    contract = AutoContractGenerator.generate_contract(headers=headers)
    for m in contract["canonical_mappings"]:
        assert "confidence" in m
        assert "detection_source" in m
        assert 0.0 <= m["confidence"] <= 1.0

# =====================================================================
# 4. Contract Schema Validation
# =====================================================================

def test_validate_contract_valid():
    valid_contract = {
        "dataset_id": "sample_valid",
        "quarantine_threshold_pct": 15.0,
        "canonical_mappings": [
            {
                "target_field": "full_name_ar",
                "group": "Identity",
                "type": "string",
                "aliases": ["الاسم"],
                "transformers": ["strip"]
            }
        ]
    }
    is_valid, errors = AutoContractGenerator.validate_contract(valid_contract)
    assert is_valid is True
    assert len(errors) == 0

def test_validate_contract_invalid():
    invalid_contract = {
        "quarantine_threshold_pct": 150.0,
        "canonical_mappings": [
            {
                "target_field": "full_name_ar",
                "group": "InvalidGroupType",
                "transformers": "not_a_list"
            },
            {
                "target_field": "full_name_ar",
                "group": "Identity",
                "transformers": []
            }
        ]
    }
    is_valid, errors = AutoContractGenerator.validate_contract(invalid_contract)
    assert is_valid is False
    assert any("dataset_id" in e for e in errors)
    assert any("quarantine_threshold_pct" in e for e in errors)
    assert any("InvalidGroupType" in e for e in errors)
    assert any("Duplicate" in e for e in errors)

# =====================================================================
# 5. Contract Diffing & Schema Auditing (ContractEngine)
# =====================================================================

def test_contract_diff_headers_compliant(fuzzy_engine):
    raw_headers = ["Phone Number", "Email Address", "University", "Custom Survey Column"]
    report = fuzzy_engine.diff_headers(raw_headers)
    
    assert report["resolved_count"] == 3
    assert report["unmapped_count"] == 1
    assert "Custom Survey Column" in report["unmapped_headers"]
    assert report["is_compliant"] is True
    assert len(report["missing_required"]) == 0

def test_contract_diff_headers_missing_required(fuzzy_engine):
    raw_headers = ["Email Address", "University"]
    report = fuzzy_engine.diff_headers(raw_headers)

    assert report["is_compliant"] is False
    assert "phone" in report["missing_required"]

# =====================================================================
# 6. Pluggable Transformer NLP Assistant Integration
# =====================================================================

class MockCustomNLPAssistant(BaseNLPAssistant):
    def classify_column(self, header: str, sample_values=None):
        if "رأي" in str(header):
            return {
                "target_field": "user_opinion",
                "group": "Free-Form",
                "type": "text",
                "transformers": ["strip"],
                "confidence": 0.88
            }
        return None

def test_pluggable_nlp_assistant_invocation():
    mock_nlp = MockCustomNLPAssistant()
    res = AutoContractGenerator.classify_column("ما هو رأيك بالمشروع؟", nlp_assistant=mock_nlp)
    assert res["target_field"] == "user_opinion"
    assert res["group"] == "Free-Form"
    assert res["detection_source"] == "transformer_nlp"

def test_rule_assisted_nlp_built_in():
    rule_nlp = RuleAssistedNLP()
    res = rule_nlp.classify_column("General Feedback & Comments")
    assert res is not None
    assert res["target_field"] == "feedback"
    assert res["group"] == "Free-Form"

    res_score = rule_nlp.classify_column("معدل التقييم")
    assert res_score is not None
    assert res_score["target_field"] == "evaluation_score"
    assert res_score["group"] == "Structured/MetaData"

def test_transformer_nlp_assistant_defaults():
    nlp_assistant = TransformerNLPAssistant()
    res = nlp_assistant.classify_column("General Feedback & Comments")
    assert res is not None
    assert res["group"] == "Free-Form"
    assert res["detection_source"] == "transformer_nlp"

def test_trained_model_multiclass_inference():
    nlp = TransformerNLPAssistant()
    assert nlp.sklearn_pipeline is not None, "Trained scikit-learn model should be loaded"

    # Identity
    id_res = nlp.classify_column("رقم الجوال", sample_values=["0551234567"])
    assert id_res["group"] == "Identity"
    assert id_res["confidence"] > 0.80

    # Structured/MetaData
    meta_res = nlp.classify_column("السنة الدراسية", sample_values=["سنة ثالثة"])
    assert meta_res["group"] == "Structured/MetaData"

    # Categorical/MultiSelect
    cat_res = nlp.classify_column("هل توافق على الشروط والأحكام؟", sample_values=["نعم", "لا"])
    assert cat_res["group"] == "Categorical/MultiSelect"

    # Free-Form
    free_res = nlp.classify_column("لماذا ترغب بالانضمام إلى هذا المعسكر؟", sample_values=["لتطوير مهاراتي في هندسة البيانات"])
    assert free_res["group"] == "Free-Form"
    assert free_res["confidence"] > 0.85

# =====================================================================
# 7. Fuzzy Boundary & Arabic Variations Stress Tests
# =====================================================================

def test_fuzzy_arabic_hamza_and_noise_variations(fuzzy_engine):
    # Test slight variations in Arabic spelling
    test_cases = [
        ("الاسم الثلاثيي", "full_name_ar"),
        ("اسم المشاارك", "full_name_ar"),
        ("رقم الجووال", "phone")
    ]
    for h, expected in test_cases:
        resolved, _ = fuzzy_engine.resolve_headers([h])
        assert resolved.get(h) == expected, f"Failed for {h}"

def test_fuzzy_custom_threshold_tuning():
    contract = {
        "dataset_id": "threshold_test",
        "quarantine_threshold_pct": 15.0,
        "similarity_threshold": 0.95, # Very strict threshold
        "canonical_mappings": [
            {"target_field": "phone", "group": "Identity", "type": "phone", "aliases": ["Phone Number"]}
        ]
    }
    strict_engine = ContractEngine(contract)
    # Typo "Phoone Nummber" has similarity ~0.92, which is < 0.95, so it should be unmapped
    resolved, unmapped = strict_engine.resolve_headers(["Phoone Nummber"])
    assert len(resolved) == 0
    assert "Phoone Nummber" in unmapped

# =====================================================================
# 8. Deep Sample Inference Stress Tests (Edge Cases & Messy Data)
# =====================================================================

def test_deep_sample_excel_float_phone_numbers():
    # Excel frequently reads numbers as floats (e.g. 501234567.0)
    samples = [501234567.0, "509876543.0", 551122334]
    res = AutoContractGenerator.classify_column("رقم_بدون_عنوان", sample_values=samples)
    assert res["target_field"] == "phone"
    assert res["group"] == "Identity"
    assert res["type"] == "phone"

def test_deep_sample_mixed_none_and_sparse_samples():
    # Mostly None or blank values, but with valid academic email
    sparse_samples = [None, "", "   ", "fahad.student@uj.edu.sa", None, ""]
    res = AutoContractGenerator.classify_column("unknown_field_12", sample_values=sparse_samples)
    assert res["target_field"] == "academic_email"
    assert res["group"] == "Identity"
    assert res["confidence"] >= 0.90

def test_deep_sample_long_arabic_free_form_text():
    samples = [
        "نأمل منكم التكرم بإتاحة الفرصة لنا للمشاركة والاستفادة من هذه التجربة القيمة ونشكركم جزيل الشكر",
        "أود تعلم خوارزميات الذكاء الاصطناعي وبناء نماذج عملية تخدم المجتمع"
    ]
    res = AutoContractGenerator.classify_column("ملاحظات_حرة", sample_values=samples)
    assert res["group"] == "Free-Form"
    assert res["type"] == "text"
    assert res["confidence"] >= 0.80

# =====================================================================
# 9. Real-World Datasets End-to-End Stress Test
# =====================================================================

def test_real_dataset_developers_committee_55_columns():
    file_path = r"c:\Users\Admin\Documents\DataCleaning\RawData\Developers_Committee_2026_Responses.xlsx"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0] if h is not None]
    
    assert len(headers) == 55, f"Expected 55 headers, got {len(headers)}"

    sample_rows = [dict(zip(headers, r)) for r in rows[1:6]]
    contract = AutoContractGenerator.generate_contract(
        headers=headers,
        sample_rows=sample_rows,
        dataset_id="committee_stress_test"
    )

    # 1. Verify contract validity
    is_valid, errors = AutoContractGenerator.validate_contract(contract)
    assert is_valid is True, f"Generated contract failed validation: {errors}"

    # 2. Verify all 55 mappings exist and all target fields are distinct
    mappings = contract["canonical_mappings"]
    assert len(mappings) == 55
    target_fields = [m["target_field"] for m in mappings]
    assert len(target_fields) == len(set(target_fields)), "Duplicate target fields found in 55-column contract!"

    # 3. Verify all 4 guideline groups are represented
    groups_found = set(m["group"] for m in mappings)
    assert "Identity" in groups_found
    assert "Structured/MetaData" in groups_found
    assert "Categorical/MultiSelect" in groups_found
    assert "Free-Form" in groups_found

    # 4. Verify contract diffing is compliant
    engine = ContractEngine(contract)
    audit = engine.diff_headers(headers)
    assert audit["resolved_count"] == len(set(headers)) # 53 unique headers across 55 columns
    assert audit["unmapped_count"] == 0
    assert audit["is_compliant"] is True

def test_real_dataset_workshop_contract_generation():
    file_path = r"c:\Users\Admin\Documents\DataCleaning\RawData\Distraction_Era_Workshop_Dataset.xlsx"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0] if h is not None]

    contract = AutoContractGenerator.generate_contract(headers=headers, dataset_id="workshop_auto")
    is_valid, errors = AutoContractGenerator.validate_contract(contract)
    assert is_valid is True

    engine = ContractEngine(contract)
    audit = engine.diff_headers(headers)
    assert audit["is_compliant"] is True
    assert audit["resolved_count"] == len(headers)

def test_real_dataset_multisheet_contract_generation():
    file_path = r"c:\Users\Admin\Documents\DataCleaning\RawData\MultiSheet_Workshop_Batches.xlsx"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        headers = [str(h) for h in rows[0] if h is not None]
        contract = AutoContractGenerator.generate_contract(headers=headers, dataset_id=f"sheet_{sheet_name}")
        is_valid, _ = AutoContractGenerator.validate_contract(contract)
        assert is_valid is True

# =====================================================================
# 10. Transformer NLP Assistant Fault-Tolerance & Resilience
# =====================================================================

class CrashingNLPAssistant(BaseNLPAssistant):
    def classify_column(self, header: str, sample_values=None):
        raise ConnectionResetError("Simulated neural inference crash / timeout!")

def test_nlp_assistant_fault_tolerance_exception_handling():
    crashing_nlp = CrashingNLPAssistant()
    res = AutoContractGenerator.classify_column("obscure_column_header", nlp_assistant=crashing_nlp)
    assert res is not None
    assert res["group"] == "Structured/MetaData"
    assert res["detection_source"] == "fallback"
    assert res["confidence"] < 0.60

def test_transformer_nlp_softmax_calibrated_confidence():
    logits = [2.0, 1.0, 0.1, -1.0]
    probs = TransformerNLPAssistant._softmax(logits)
    assert len(probs) == 4
    assert sum(probs) == pytest.approx(1.0, abs=1e-5)
    assert probs[0] > probs[1] > probs[2] > probs[3]

# =====================================================================
# 11. Standalone CLI Execution Tests
# =====================================================================

def test_cli_validate_success():
    contract_file = r"c:\Users\Admin\Documents\DataCleaning\config\workshop_contract.yaml"
    with unittest.mock.patch("sys.argv", ["contract_generator", "--validate", contract_file]):
        # Must execute without error
        cli_main()

def test_cli_validate_failure():
    with tempfile.NamedTemporaryFile(suffix=".yaml", delete=False, mode="w", encoding="utf-8") as f:
        f.write("quarantine_threshold_pct: 120.0\ncanonical_mappings: []\n")
        bad_contract = f.name
    try:
        with unittest.mock.patch("sys.argv", ["contract_generator", "--validate", bad_contract]):
            with pytest.raises(SystemExit) as excinfo:
                cli_main()
            assert excinfo.value.code == 1
    finally:
        if os.path.exists(bad_contract):
            os.remove(bad_contract)

def test_cli_diff():
    contract_file = r"c:\Users\Admin\Documents\DataCleaning\config\workshop_contract.yaml"
    data_file = r"c:\Users\Admin\Documents\DataCleaning\RawData\Distraction_Era_Workshop_Dataset.xlsx"
    with unittest.mock.patch("sys.argv", ["contract_generator", "--diff", contract_file, "--input", data_file]):
        cli_main()

def test_cli_generate():
    data_file = r"c:\Users\Admin\Documents\DataCleaning\RawData\Distraction_Era_Workshop_Dataset.xlsx"
    with tempfile.TemporaryDirectory() as tmpdir:
        out_yaml = os.path.join(tmpdir, "cli_generated_contract.yaml")
        with unittest.mock.patch("sys.argv", ["contract_generator", "--input", data_file, "--output", out_yaml]):
            cli_main()
        assert os.path.exists(out_yaml)
        with open(out_yaml, "r", encoding="utf-8") as f:
            loaded = yaml.safe_load(f)
        assert len(loaded["canonical_mappings"]) > 0
        valid, _ = AutoContractGenerator.validate_contract(loaded)
        assert valid is True
