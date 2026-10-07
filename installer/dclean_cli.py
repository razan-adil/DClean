"""
DClean Global Command-Line Interface (CLI).

Provides unified command dispatch for cleaning, contract validation,
schema diffing, automated contract generation, and GUI invocation.
"""

import sys
import os
import argparse

# Ensure project root is in sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

def cmd_clean(args):
    from pipeline.runner import PipelineRunner
    input_file = os.path.abspath(args.input)
    if not os.path.exists(input_file):
        print(f"[Error] Input file '{input_file}' does not exist.")
        sys.exit(1)

    contract = os.path.abspath(args.contract) if args.contract else None
    out_excel = os.path.abspath(args.output) if args.output else None
    db_path = os.path.abspath(args.db) if args.db else "CleanData/cleaned_database.db"

    print("=" * 70)
    print("   DCLEAN DATA CLEANING & ENTITY RESOLUTION PIPELINE")
    print("=" * 70)
    print(f"Input:       {input_file}")
    print(f"Contract:    {contract or '[Auto-Detect via Local NLP Sequence Classifier]'}")
    print(f"SQLite DB:   {db_path}")
    print("-" * 70)

    runner = PipelineRunner(contract_path=contract, db_path=db_path)
    summary = runner.run(input_file_path=input_file, output_excel_path=out_excel)

    out_file = out_excel or f"{os.path.splitext(input_file)[0]}_Cleaned.xlsx"
    print("\n" + "=" * 70)
    print("   EXECUTION & GOVERNANCE AUDIT SUMMARY")
    print("=" * 70)
    print(f"  Source File:                {summary.source_file}")
    print(f"  Processed Sheet(s):         {summary.sheet_name}")
    print(f"  Total Ingested Rows:        {summary.total_rows}")
    print(f"  Valid Processed Rows:       {summary.valid_rows}")
    print(f"  Quarantined Rows:           {summary.quarantined_rows} ({summary.quarantine_rate_pct}%)")
    print(f"  Unique Entities Resolved:   {summary.unique_persons_resolved}")
    print(f"  Duplicate Submissions Merged:{summary.duplicate_events_merged}")
    print(f"  Pipeline Health Status:     [{summary.status}]")
    print("-" * 70)
    print(f"  Clean Database:             file:///{db_path.replace(os.sep, '/')}")
    print(f"  Clean Excel Output:         file:///{out_file.replace(os.sep, '/')}")
    print("=" * 70 + "\n")

def cmd_validate(args):
    import yaml
    from pipeline.contract_generator import AutoContractGenerator
    contract_path = os.path.abspath(args.contract)
    if not os.path.exists(contract_path):
        print(f"[ERROR] Contract file '{contract_path}' not found.")
        sys.exit(1)
    with open(contract_path, "r", encoding="utf-8") as f:
        c_dict = yaml.safe_load(f)
    valid, errors = AutoContractGenerator.validate_contract(c_dict)
    if valid:
        print(f"[VALID] Contract '{args.contract}' is valid and guideline compliant.")
    else:
        print(f"[INVALID] Contract '{args.contract}' has errors:")
        for err in errors:
            print(f"  - {err}")
        sys.exit(1)

def cmd_diff(args):
    import openpyxl
    import csv
    from pipeline.contract_engine import ContractEngine
    contract_path = os.path.abspath(args.contract)
    input_path = os.path.abspath(args.input)
    if not os.path.exists(contract_path) or not os.path.exists(input_path):
        print("[ERROR] Input file or contract does not exist.")
        sys.exit(1)
    engine = ContractEngine(contract_path)
    ext = os.path.splitext(input_path)[1].lower()
    if ext in [".xlsx", ".xlsm"]:
        wb = openpyxl.load_workbook(input_path, data_only=True)
        ws = wb.active
        headers = [str(c.value) for c in next(ws.iter_rows(values_only=False)) if c.value is not None]
    else:
        with open(input_path, "r", encoding="utf-8", errors="replace") as f:
            headers = next(csv.reader(f), [])
    res = engine.diff_headers(headers)
    print(f"\nAudit Diff for '{os.path.basename(input_path)}' vs '{os.path.basename(contract_path)}':")
    print(f"  Resolved:    {res['resolved_count']} / {len(headers)} columns")
    print(f"  Unmapped:    {res['unmapped_count']}")
    print(f"  Compliant:   {res['is_compliant']}")
    if res['missing_required']:
        print(f"  Missing:     {res['missing_required']}\n")

def cmd_generate(args):
    import openpyxl
    import csv
    from pipeline.contract_generator import AutoContractGenerator
    in_path = os.path.abspath(args.input)
    if not os.path.exists(in_path):
        print(f"[Error] Input file '{in_path}' not found.")
        sys.exit(1)
    dataset_id = os.path.splitext(os.path.basename(in_path))[0]
    out_yaml = args.output or f"config/{dataset_id}_contract.yaml"

    ext = os.path.splitext(in_path)[1].lower()
    rows = []
    if ext in [".xlsx", ".xlsm"]:
        wb = openpyxl.load_workbook(in_path, data_only=True)
        ws = wb.active
        iter_rows = ws.iter_rows(values_only=True)
        raw_headers = [str(h) for h in next(iter_rows, []) if h is not None]
        for r in iter_rows:
            if any(v is not None for v in r):
                rows.append(dict(zip(raw_headers, r)))
            if len(rows) >= 5:
                break
    else:
        with open(in_path, "r", encoding="utf-8", errors="replace") as f:
            r = csv.DictReader(f)
            raw_headers = r.fieldnames or []
            for row in r:
                rows.append(row)
                if len(rows) >= 5:
                    break

    AutoContractGenerator.generate_contract(
        headers=raw_headers,
        sample_rows=rows,
        dataset_id=dataset_id,
        output_yaml_path=out_yaml
    )
    print(f"[SUCCESS] Generated declarative contract saved to: {out_yaml}")

def cmd_gui(args):
    from installer.dclean_launcher import launch_desktop_gui
    launch_desktop_gui()

def cmd_dashboard(args):
    from installer.dclean_launcher import launch_dashboard
    if not launch_dashboard():
        print("[Notice] Dashboard not uploaded yet. Launching Desktop GUI instead...")
        from installer.dclean_launcher import launch_desktop_gui
        launch_desktop_gui()

def main():
    parser = argparse.ArgumentParser(
        prog="dclean",
        description="DClean: Production Data Cleaning & Entity Resolution System"
    )
    subparsers = parser.add_subparsers(dest="command", help="Available commands")

    # clean command
    p_clean = subparsers.add_parser("clean", help="Clean and deduplicate tabular dataset (.xlsx / .csv)")
    p_clean.add_argument("--input", "-i", required=True, help="Path to input raw dataset")
    p_clean.add_argument("--contract", "-c", default=None, help="Path to YAML contract (auto-detected if omitted)")
    p_clean.add_argument("--output", "-o", default=None, help="Path to output cleaned Excel workbook")
    p_clean.add_argument("--db", "-d", default=None, help="Path to SQLite database")

    # validate command
    p_val = subparsers.add_parser("validate", help="Validate YAML schema contract")
    p_val.add_argument("--contract", "-c", required=True, help="Path to YAML contract")

    # diff command
    p_diff = subparsers.add_parser("diff", help="Diff dataset headers against contract")
    p_diff.add_argument("--input", "-i", required=True, help="Path to tabular input file")
    p_diff.add_argument("--contract", "-c", required=True, help="Path to YAML contract")

    # generate command
    p_gen = subparsers.add_parser("generate", help="Auto-generate YAML contract from dataset")
    p_gen.add_argument("--input", "-i", required=True, help="Path to input tabular file")
    p_gen.add_argument("--output", "-o", default=None, help="Output YAML file path")

    # gui command
    subparsers.add_parser("gui", help="Open the desktop GUI window")

    # dashboard command
    subparsers.add_parser("dashboard", help="Launch Streamlit dashboard")

    # If no arguments provided, display help
    if len(sys.argv) == 1:
        parser.print_help()
        return

    args = parser.parse_args()
    if args.command == "clean":
        cmd_clean(args)
    elif args.command == "validate":
        cmd_validate(args)
    elif args.command == "diff":
        cmd_diff(args)
    elif args.command == "generate":
        cmd_generate(args)
    elif args.command == "gui":
        cmd_gui(args)
    elif args.command == "dashboard":
        cmd_dashboard(args)
    else:
        parser.print_help()

if __name__ == "__main__":
    main()
